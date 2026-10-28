
from datetime import date, datetime
from io import BytesIO

from docx import Document
from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy import or_, select

from ..deps import DB, AthletePrincipal, CoachPrincipal, CurrentPrincipal
from ..models import (
    AthleteProfile,
    AthleteRank,
    CoachAthleteLink,
    CoachLinkEvent,
    Consent,
    ConsentKind,
    Discipline,
    Gender,
    LinkStatus,
    Role,
    TrainingStage,
    User,
)
from ..services.audit import log_action
from ..services.notifications import notify
from ..services.telegram import link_url

router = APIRouter(prefix="/api", tags=["profile"])

CONSENT_TEXT = (
    "Я, {signer}, в соответствии с Федеральным законом от 27.07.2006 № 152-ФЗ «О персональных данных» "
    "даю согласие региональному отделению Федерации спортивного программирования Оренбургской области "
    "на обработку персональных данных{subject}: фамилия, имя, отчество, дата рождения, муниципалитет, "
    "организация, контактные данные, сведения о спортивной подготовке, результатах и разрядах - в целях учёта "
    "занимающихся, организации соревнований и подготовки отчётных документов. Согласие действует до его отзыва."
)


class AthleteBlock(BaseModel):
    organization_id: int | None
    start_year: int
    training_stage: TrainingStage | None
    discipline_ids: list[int]
    guardian_name: str | None
    guardian_phone: str | None
    rusada_course_date: date | None
    rusada_certificate_url: str | None
    coaches: list[dict]
    ranks: list[dict]


class ConsentOut(BaseModel):
    id: int
    kind: ConsentKind
    signed_by: str
    signed_at: datetime
    text_version: str
    revoked_at: datetime | None


class ProfileOut(BaseModel):
    id: int
    email: str
    last_name: str
    first_name: str
    middle_name: str | None
    birth_date: date | None
    gender: Gender | None
    phone: str | None
    municipality_id: int | None
    roles: list[Role]
    is_minor: bool
    telegram_linked: bool
    telegram_link: str | None
    athlete: AthleteBlock | None
    consents: list[ConsentOut]


def _is_minor(u: User) -> bool:
    if not u.birth_date:
        return False
    t = date.today()
    return t.year - u.birth_date.year - ((t.month, t.day) < (u.birth_date.month, u.birth_date.day)) < 18


def profile_out(db: DB, u: User) -> ProfileOut:
    athlete = None
    if u.athlete:
        a = u.athlete
        links = db.scalars(select(CoachAthleteLink).where(CoachAthleteLink.athlete_id == u.id)).all()
        ranks = db.scalars(select(AthleteRank).where(AthleteRank.athlete_id == u.id)).all()
        athlete = AthleteBlock(
            organization_id=a.organization_id,
            start_year=a.start_year,
            training_stage=a.training_stage,
            discipline_ids=[d.id for d in a.disciplines],
            guardian_name=a.guardian_name,
            guardian_phone=a.guardian_phone,
            rusada_course_date=a.rusada_course_date,
            rusada_certificate_url=a.rusada_certificate_url,
            coaches=[
                {"link_id": l.id, "coach": l.coach.full_name, "discipline": l.discipline.name, "status": l.status.value}
                for l in links
            ],
            ranks=[
                {"discipline": r.discipline.name, "rank": r.rank.name, "valid_until": r.valid_until.isoformat() if r.valid_until else None}
                for r in ranks
            ],
        )
    consents = db.scalars(select(Consent).where(Consent.user_id == u.id).order_by(Consent.signed_at)).all()
    return ProfileOut(
        id=u.id,
        email=u.email,
        last_name=u.last_name,
        first_name=u.first_name,
        middle_name=u.middle_name,
        birth_date=u.birth_date,
        gender=u.gender,
        phone=u.phone,
        municipality_id=u.municipality_id,
        roles=sorted(u.role_set),
        is_minor=_is_minor(u),
        telegram_linked=bool(u.telegram_chat_id),
        telegram_link=link_url(u.id),
        athlete=athlete,
        consents=[ConsentOut.model_validate(c, from_attributes=True) for c in consents],
    )


@router.get("/me/profile", response_model=ProfileOut)
def get_profile(db: DB, p: CurrentPrincipal):
    return profile_out(db, p.user)


class ProfileIn(BaseModel):
    last_name: str = Field(min_length=1)
    first_name: str = Field(min_length=1)
    middle_name: str | None = None
    phone: str | None = None
    municipality_id: int | None = None
    organization_id: int | None = None
    discipline_ids: list[int] | None = None
    start_year: int | None = None
    guardian_name: str | None = None
    guardian_phone: str | None = None


@router.put("/me/profile", response_model=ProfileOut)
def update_profile(data: ProfileIn, db: DB, p: CurrentPrincipal):
    u = p.user
    for f in ("last_name", "first_name", "middle_name", "phone", "municipality_id"):
        setattr(u, f, getattr(data, f))
    if u.athlete:
        a = u.athlete
        a.organization_id = data.organization_id
        if data.discipline_ids:
            a.disciplines = list(db.scalars(select(Discipline).where(Discipline.id.in_(data.discipline_ids))))
        if data.start_year:
            a.start_year = data.start_year
        if _is_minor(u) and not data.guardian_name:
            raise HTTPException(422, "Для несовершеннолетнего укажите законного представителя")
        a.guardian_name = data.guardian_name
        a.guardian_phone = data.guardian_phone
    log_action(db, p, "update", "profile", u.id)
    db.commit()
    return profile_out(db, u)


class AntidopingIn(BaseModel):
    rusada_course_date: date | None
    rusada_certificate_url: str | None = None


@router.put("/me/antidoping", response_model=ProfileOut)
def update_antidoping(data: AntidopingIn, db: DB, p: AthletePrincipal):
    a = p.user.athlete
    if a is None:
        raise HTTPException(404, "Профиль спортсмена не найден")
    if data.rusada_course_date and data.rusada_course_date > date.today():
        raise HTTPException(422, "Дата прохождения курса не может быть в будущем")
    a.rusada_course_date = data.rusada_course_date
    a.rusada_certificate_url = data.rusada_certificate_url
    log_action(db, p, "update", "antidoping", p.id)
    db.commit()
    return profile_out(db, p.user)


@router.get("/me/consents/{consent_id}.docx")
def consent_document(consent_id: int, db: DB, p: CurrentPrincipal):
    c = db.get(Consent, consent_id)
    if c is None or c.user_id != p.id:
        raise HTTPException(404, "Согласие не найдено")
    doc = Document()
    doc.add_heading("Согласие на обработку персональных данных", level=1)
    subject = f" несовершеннолетнего {p.user.full_name}" if c.kind == ConsentKind.PD_GUARDIAN else ""
    doc.add_paragraph(CONSENT_TEXT.format(signer=c.signed_by, subject=subject))
    doc.add_paragraph(f"Дата: {c.signed_at:%d.%m.%Y}. Редакция текста: {c.text_version}. Подписано электронно на платформе.")
    buf = BytesIO()
    doc.save(buf)
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="consent_{c.id}.docx"'},
    )




class AthleteSearchOut(BaseModel):
    id: int
    full_name: str
    birth_year: int | None
    organization: str | None
    disciplines: list[dict]


@router.get("/coach/search", response_model=list[AthleteSearchOut])
def search_athletes(q: str, db: DB, _: CoachPrincipal):
    if len(q.strip()) < 2:
        return []
    like = f"%{q.strip()}%"
    profiles = db.scalars(
        select(AthleteProfile)
        .join(User, AthleteProfile.user_id == User.id)
        .where(or_(User.last_name.ilike(like), User.first_name.ilike(like)))
        .order_by(User.last_name)
        .limit(20)
    ).all()
    return [
        AthleteSearchOut(
            id=a.user_id,
            full_name=a.user.full_name,
            birth_year=a.user.birth_date.year if a.user.birth_date else None,
            organization=(a.organization.short_name or a.organization.name) if a.organization else "без организации",
            disciplines=[{"id": d.id, "name": d.name} for d in a.disciplines],
        )
        for a in profiles
    ]


class LinkIn(BaseModel):
    athlete_id: int
    discipline_id: int


class LinkOut(BaseModel):
    id: int
    coach_id: int
    coach_name: str
    athlete_id: int
    athlete_name: str
    discipline: str
    status: LinkStatus
    events: list[dict]


def link_out(l: CoachAthleteLink) -> LinkOut:
    return LinkOut(
        id=l.id,
        coach_id=l.coach_id,
        coach_name=l.coach.full_name,
        athlete_id=l.athlete_id,
        athlete_name=l.athlete.full_name,
        discipline=l.discipline.name,
        status=l.status,
        events=[{"status": e.status.value, "at": e.at.isoformat(), "comment": e.comment} for e in l.events],
    )


@router.post("/coach/links", response_model=LinkOut, status_code=201)
def request_link(data: LinkIn, db: DB, p: CoachPrincipal):
    athlete = db.get(AthleteProfile, data.athlete_id)
    if athlete is None:
        raise HTTPException(404, "Спортсмен не найден")
    if data.discipline_id not in {d.id for d in athlete.disciplines}:
        raise HTTPException(422, "Спортсмен не занимается этой дисциплиной")
    if data.athlete_id == p.id:
        raise HTTPException(422, "Нельзя привязать себя")
    link = db.scalar(
        select(CoachAthleteLink).where(
            CoachAthleteLink.coach_id == p.id,
            CoachAthleteLink.athlete_id == data.athlete_id,
            CoachAthleteLink.discipline_id == data.discipline_id,
        )
    )
    if link and link.status in (LinkStatus.PENDING, LinkStatus.ACTIVE, LinkStatus.SUSPENDED):
        raise HTTPException(409, "Связь уже существует")
    link = link or CoachAthleteLink(coach_id=p.id, athlete_id=data.athlete_id, discipline_id=data.discipline_id)
    link.status = LinkStatus.PENDING
    link.events.append(CoachLinkEvent(status=LinkStatus.PENDING, by_user_id=p.id, comment="Запрос тренера"))
    db.add(link)
    db.flush()
    notify(db, data.athlete_id, "Запрос от тренера", f"{p.user.full_name} хочет стать вашим тренером", "/athlete/profile")
    log_action(db, p, "create", "coach_link", link.id)
    db.commit()
    return link_out(link)


@router.get("/me/links", response_model=list[LinkOut])
def my_links(db: DB, p: CurrentPrincipal):
    col = CoachAthleteLink.athlete_id if p.role == Role.ATHLETE else CoachAthleteLink.coach_id
    return [link_out(l) for l in db.scalars(select(CoachAthleteLink).where(col == p.id).order_by(CoachAthleteLink.id))]


class LinkStatusIn(BaseModel):
    status: LinkStatus
    comment: str | None = None


ALLOWED = {
    LinkStatus.PENDING: {LinkStatus.ACTIVE, LinkStatus.DECLINED},
    LinkStatus.ACTIVE: {LinkStatus.SUSPENDED, LinkStatus.ENDED},
    LinkStatus.SUSPENDED: {LinkStatus.ACTIVE, LinkStatus.ENDED},
}


@router.post("/links/{link_id}/status", response_model=LinkOut)
def change_link(link_id: int, data: LinkStatusIn, db: DB, p: CurrentPrincipal):
    link = db.get(CoachAthleteLink, link_id)
    if link is None:
        raise HTTPException(404, "Связь не найдена")
    if p.id not in (link.coach_id, link.athlete_id) and p.role != Role.ADMIN:
        raise HTTPException(403, "Нет доступа к этой связи")
    if data.status not in ALLOWED.get(link.status, set()):
        raise HTTPException(409, "Недопустимая смена статуса")
    if link.status == LinkStatus.PENDING and p.id != link.athlete_id:
        raise HTTPException(403, "Запрос подтверждает спортсмен")
    link.status = data.status
    link.events.append(CoachLinkEvent(status=data.status, by_user_id=p.id, comment=data.comment))
    other = link.coach_id if p.id == link.athlete_id else link.athlete_id
    notify(db, other, "Связь тренер–спортсмен", f"{link.discipline.name}: статус «{data.status.value}»", None)
    log_action(db, p, "status", "coach_link", link.id, new=data.status.value)
    db.commit()
    return link_out(link)

