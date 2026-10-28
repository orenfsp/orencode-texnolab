
import secrets
from datetime import date
from io import BytesIO
from typing import Annotated

import segno
from docx.shared import Cm
from fastapi import APIRouter, BackgroundTasks, Body, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from pydantic import BaseModel
from sqlalchemy import delete, func, select

from ..config import settings
from ..deps import DB, AdminPrincipal, AthletePrincipal, CurrentPrincipal, ReaderPrincipal
from ..models import (
    Appeal,
    AppealStatus,
    AthleteProfile,
    CaseEvent,
    CoachAthleteLink,
    Competition,
    CompetitionStatus,
    Disqualification,
    EvskNorm,
    Evidence,
    JudgeAssignment,
    LinkStatus,
    Organization,
    ParticipationCertificate,
    PetitionStatus,
    PrizePetition,
    PrizeRule,
    RatingLevel,
    Result,
    Role,
    SportRank,
    User,
    Violation,
    ViolationDecision,
)
from ..models.base import utcnow
from ..services import documents as docs
from ..services import evsk
from ..services.audit import log_action
from ..services.competitions import check_eligibility
from ..services.notifications import deliver_in_background, notify, notify_role
from ..services.protocol import assign_places
from ..services.rating import LEVEL_RU
from ..services.storage import ALLOWED_EVIDENCE, folder, save_upload

router = APIRouter(prefix="/api", tags=["extras"])




class CertificateIn(BaseModel):
    result_id: int
    purpose: str | None = None


class CertificateOut(BaseModel):
    id: int
    number: str
    verify_code: str
    verify_url: str
    competition: str
    competition_date: date
    discipline: str
    place: int | None
    participants: int
    athlete_name: str
    purpose: str | None
    issued_at: str


def _verify_url(code: str) -> str:
    return f"{settings.public_base_url.rstrip('/')}/verify/{code}"


def cert_out(db: DB, c: ParticipationCertificate) -> CertificateOut:
    res = c.result
    comp = res.competition
    return CertificateOut(
        id=c.id, number=c.number, verify_code=c.verify_code, verify_url=_verify_url(c.verify_code),
        competition=comp.title, competition_date=comp.start_date, discipline=comp.discipline.name,
        place=res.place,
        participants=db.scalar(select(func.count(Result.id)).where(Result.competition_id == comp.id, Result.place.is_not(None))),
        athlete_name=c.athlete.full_name, purpose=c.purpose, issued_at=c.issued_at.isoformat(),
    )


@router.post("/me/certificates", response_model=CertificateOut, status_code=201)
def issue_certificate(data: CertificateIn, db: DB, p: AthletePrincipal):
    res = db.get(Result, data.result_id)
    if res is None or res.athlete_id != p.id or res.competition.status != CompetitionStatus.COMPLETED:
        raise HTTPException(422, "Справку можно получить только по своему результату из опубликованного протокола")
    c = ParticipationCertificate(
        number="", verify_code=secrets.token_urlsafe(12), athlete_id=p.id, result_id=res.id, purpose=data.purpose,
    )
    db.add(c)
    db.flush()
    c.number = f"СП-{date.today().year}-{c.id:06d}"
    log_action(db, p, "create", "certificate", c.id, number=c.number)
    db.commit()
    return cert_out(db, c)


@router.get("/me/certificates", response_model=list[CertificateOut])
def my_certificates(db: DB, p: AthletePrincipal):
    return [cert_out(db, c) for c in db.scalars(
        select(ParticipationCertificate).where(ParticipationCertificate.athlete_id == p.id).order_by(ParticipationCertificate.id.desc())
    )]


@router.get("/certificates/{cert_id}.docx")
def certificate_docx(cert_id: int, db: DB, p: CurrentPrincipal):
    c = db.get(ParticipationCertificate, cert_id)
    if c is None or (c.athlete_id != p.id and p.role != Role.ADMIN):
        raise HTTPException(404, "Справка не найдена")
    o = cert_out(db, c)
    u = c.athlete
    doc = docs.new_doc(f"СПРАВКА № {o.number}", [])
    doc.add_paragraph(
        f"Дана {u.full_name}"
        + (f", {u.birth_date:%d.%m.%Y} года рождения," if u.birth_date else ",")
        + f" в том, что он(а) принял(а) участие в соревновании «{o.competition}» "
        f"(дисциплина «{o.discipline}»), проведённом {o.competition_date:%d.%m.%Y}, "
        + (f"и занял(а) {o.place} место из {o.participants} участников." if o.place else "без итогового места.")
    )
    if o.purpose:
        doc.add_paragraph(f"Справка выдана для предъявления: {o.purpose}.")
    doc.add_paragraph(f"Дата выдачи: {c.issued_at:%d.%m.%Y}. Подлинность справки можно проверить по QR-коду или по адресу:")
    doc.add_paragraph(o.verify_url)
    buf = BytesIO()
    segno.make(o.verify_url, error="m").save(buf, kind="png", scale=6, border=2)
    buf.seek(0)
    doc.add_picture(buf, width=Cm(4))
    docs.signature(doc, "Председатель регионального отделения")
    return docs.file_response(docs.docx_bytes(doc), f"certificate_{o.number}.docx")


@router.get("/certificates/{cert_id}/qr.svg")
def certificate_qr(cert_id: int, db: DB, p: CurrentPrincipal):
    from fastapi.responses import Response

    c = db.get(ParticipationCertificate, cert_id)
    if c is None or (c.athlete_id != p.id and p.role != Role.ADMIN):
        raise HTTPException(404, "Справка не найдена")
    buf = BytesIO()
    segno.make(_verify_url(c.verify_code), error="m").save(buf, kind="svg", scale=6, border=2, xmldecl=False)
    return Response(buf.getvalue(), media_type="image/svg+xml")


@router.get("/public/certificates/{code}")
def verify_certificate(code: str, db: DB):
    c = db.scalar(select(ParticipationCertificate).where(ParticipationCertificate.verify_code == code))
    if c is None:
        raise HTTPException(404, "Справка с таким кодом не выдавалась")
    o = cert_out(db, c)
    return {"valid": True, **o.model_dump(exclude={"verify_code", "id"})}



DISCIPLINE_INFO = {
    "algo": "Олимпиадные задачи на алгоритмы: решение за ограниченное время, проверка автоматической системой.",
    "product": "Командная разработка цифрового продукта по кейсу: от идеи до работающего прототипа и защиты.",
    "infosec": "Соревнования по защите и атаке информационных систем (CTF): поиск уязвимостей, расследования.",
    "ai": "Задачи машинного обучения и анализа данных: модели оцениваются по качеству на скрытых данных.",
    "robotics": "Программирование роботов для прохождения трасс и выполнения заданий на полигоне.",
    "uas": "Программирование беспилотных летательных аппаратов: автономные полёты и миссии.",
}


@router.get("/public/overview")
def public_overview(db: DB):
    today = date.today()
    from ..models import Discipline

    upcoming = db.scalars(
        select(Competition)
        .where(Competition.status.in_([CompetitionStatus.REGISTRATION_OPEN, CompetitionStatus.REGISTRATION_CLOSED]),
               Competition.end_date >= today)
        .order_by(Competition.start_date).limit(6)
    ).all()
    latest = db.scalars(
        select(Competition).where(Competition.status == CompetitionStatus.COMPLETED).order_by(Competition.start_date.desc()).limit(3)
    ).all()
    orgs = db.scalars(select(Organization).where(Organization.conducts_sport_training).order_by(Organization.name)).all()
    return {
        "stats": {
            "athletes": db.scalar(select(func.count(AthleteProfile.user_id))),
            "competitions_this_year": db.scalar(select(func.count(Competition.id)).where(
                Competition.start_date >= date(today.year, 1, 1), Competition.status != CompetitionStatus.DRAFT)),
            "organizations": len(orgs),
        },
        "disciplines": [
            {"id": d.id, "code": d.code, "name": d.name, "about": DISCIPLINE_INFO.get(d.code, "")}
            for d in db.scalars(select(Discipline).order_by(Discipline.sort))
        ],
        "organizations": [
            {"id": o.id, "name": o.name, "municipality": o.municipality.name, "address": o.address,
             "disciplines": sorted({p.discipline.name for p in o.programs})}
            for o in orgs
        ],
        "upcoming": [
            {"id": c.id, "title": c.title, "discipline": c.discipline.name, "start_date": c.start_date,
             "registration_deadline": c.registration_deadline, "age_group": c.age_group.name, "format": c.format.value,
             "status": c.status.value}
            for c in upcoming
        ],
        "podiums": [
            {"competition_id": c.id, "title": c.title, "date": c.start_date, "discipline": c.discipline.name,
             "top": [{"place": r.place, "name": f"{r.athlete.last_name} {r.athlete.first_name[0]}."}
                     for r in sorted((r for r in c.results if r.place and r.place <= 3), key=lambda r: r.place)]}
            for c in latest
        ],
    }




@router.get("/me/path")
def my_path(db: DB, p: AthletePrincipal):
    u = p.user
    prof = u.athlete
    ranks = db.scalars(select(SportRank).order_by(SportRank.level)).all()
    norms = db.scalars(select(EvskNorm)).all()
    today = date.today()
    upcoming = db.scalars(
        select(Competition).where(
            Competition.status.in_([CompetitionStatus.REGISTRATION_OPEN, CompetitionStatus.REGISTRATION_CLOSED]),
            Competition.start_date >= today,
        ).order_by(Competition.start_date)
    ).all()
    out = []
    for d in prof.disciplines if prof else []:
        cur = evsk.current_rank(db, u.id, d.id)
        cur_level = cur.rank.level if cur else 0
        nxt = next((r for r in ranks if r.level > cur_level), None)
        reqs = sorted(
            [n for n in norms if nxt and n.rank_id == nxt.id],
            key=lambda n: list(RatingLevel).index(n.rating_level),
        )
        comps = []
        for c in upcoming:
            if c.discipline_id != d.id:
                continue
            ok, note = check_eligibility(c, u)
            gives = evsk.best_rank_for(db, c.rating_level, 1, 1000)
            comps.append({
                "id": c.id, "title": c.title, "start_date": c.start_date, "level": LEVEL_RU[c.rating_level],
                "eligible": ok, "note": note, "status": c.status.value,
                "can_reach_next": bool(nxt and gives and gives.level >= nxt.level),
            })
        out.append({
            "discipline": d.name,
            "current_rank": cur.rank.name if cur else None,
            "valid_until": cur.valid_until if cur else None,
            "next_rank": nxt.name if nxt else None,
            "requirements": [
                f"Занять место не ниже {n.max_place}-го на соревнованиях уровня «{LEVEL_RU[n.rating_level]}»"
                + (f" (при участии не менее {n.min_participants} спортсменов)" if n.min_participants > 1 else "")
                for n in reqs
            ],
            "competitions": comps,
        })
    return {
        "full_name": u.full_name,
        "training_stage": prof.training_stage.value if prof and prof.training_stage else None,
        "antidoping_done": bool(prof and prof.rusada_course_date),
        "disciplines": out,
    }




def _event(db: DB, subject: str, sid: int, status: str, by: int | None, comment: str | None = None) -> None:
    db.add(CaseEvent(subject=subject, subject_id=sid, status=status, by_user_id=by, comment=comment))


async def _attach(db: DB, subject: str, sid: int, files: list[UploadFile], by: int) -> None:
    for f in files:
        if not f.filename:
            continue
        name, stored = await save_upload("evidence", f, ALLOWED_EVIDENCE)
        db.add(Evidence(subject=subject, subject_id=sid, filename=name, stored_as=stored, uploaded_by_id=by))


def _case_details(db: DB, subject: str, sid: int) -> dict:
    ev = db.scalars(select(Evidence).where(Evidence.subject == subject, Evidence.subject_id == sid)).all()
    hist = db.scalars(select(CaseEvent).where(CaseEvent.subject == subject, CaseEvent.subject_id == sid).order_by(CaseEvent.at)).all()
    return {
        "evidence": [{"id": e.id, "filename": e.filename} for e in ev],
        "history": [
            {"status": h.status, "comment": h.comment, "at": h.at.isoformat(),
             "by": db.get(User, h.by_user_id).full_name if h.by_user_id else None}
            for h in hist
        ],
    }


def appeal_out(db: DB, a: Appeal) -> dict:
    res = db.scalar(select(Result).where(Result.competition_id == a.competition_id, Result.athlete_id == a.athlete_id))
    return {
        "id": a.id, "competition_id": a.competition_id, "competition": a.competition.title,
        "athlete_id": a.athlete_id, "athlete": a.athlete.full_name,
        "submitted_by": db.get(User, a.submitted_by_id).full_name, "text": a.text, "status": a.status.value,
        "decision": a.decision, "new_score": a.new_score, "current_score": res.score if res else None,
        "current_place": res.place if res else None, "created_at": a.created_at.isoformat(),
        **_case_details(db, "appeal", a.id),
    }


def _may_act_for(db: DB, p, athlete_id: int, discipline_id: int) -> bool:
    if p.role == Role.ATHLETE:
        return p.id == athlete_id
    if p.role == Role.COACH:
        return bool(db.scalar(select(CoachAthleteLink.id).where(
            CoachAthleteLink.coach_id == p.id, CoachAthleteLink.athlete_id == athlete_id,
            CoachAthleteLink.discipline_id == discipline_id, CoachAthleteLink.status == LinkStatus.ACTIVE)))
    return False


@router.post("/appeals", status_code=201)
async def submit_appeal(
    db: DB, p: CurrentPrincipal, bg: BackgroundTasks,
    competition_id: int = Form(...), text: str = Form(..., min_length=10), athlete_id: int | None = Form(None),
    files: list[UploadFile] = File(default=[]),
):
    comp = db.get(Competition, competition_id)
    if comp is None:
        raise HTTPException(404, "Соревнование не найдено")
    athlete_id = athlete_id or p.id
    if not _may_act_for(db, p, athlete_id, comp.discipline_id):
        raise HTTPException(403, "Апелляцию подаёт спортсмен или его тренер")
    if not db.scalar(select(Result.id).where(Result.competition_id == comp.id, Result.athlete_id == athlete_id)):
        raise HTTPException(422, "У спортсмена нет результата на этом соревновании")
    a = Appeal(competition_id=comp.id, athlete_id=athlete_id, submitted_by_id=p.id, text=text)
    db.add(a)
    db.flush()
    await _attach(db, "appeal", a.id, files, p.id)
    _event(db, "appeal", a.id, AppealStatus.SUBMITTED.value, p.id, "Апелляция подана")
    notify_role(db, Role.ADMIN, "Новая апелляция", f"«{comp.title}»: {a.athlete.full_name}", "/admin/appeals")
    log_action(db, p, "create", "appeal", a.id)
    db.commit()
    bg.add_task(deliver_in_background)
    return appeal_out(db, a)


@router.get("/appeals")
def list_appeals(db: DB, p: CurrentPrincipal):
    stmt = select(Appeal).order_by(Appeal.id.desc())
    if p.role == Role.ATHLETE:
        stmt = stmt.where(Appeal.athlete_id == p.id)
    elif p.role == Role.COACH:
        stmt = stmt.where(Appeal.submitted_by_id == p.id)
    elif p.role not in (Role.ADMIN, Role.MINISTRY):
        raise HTTPException(403, "Нет доступа")
    return [appeal_out(db, a) for a in db.scalars(stmt)]


class AppealDecisionIn(BaseModel):
    status: AppealStatus
    decision: str | None = None
    new_score: float | None = None


def _recalc_and_notify(db: DB, comp: Competition, reason: str) -> None:
    assign_places(comp.results)
    for r in comp.results:
        notify(db, r.athlete_id, "Протокол скорректирован", f"«{comp.title}»: {reason}", f"/protocols/{comp.id}")


@router.post("/appeals/{appeal_id}/decision")
def decide_appeal(appeal_id: int, data: AppealDecisionIn, db: DB, p: AdminPrincipal, bg: BackgroundTasks):
    a = db.get(Appeal, appeal_id)
    if a is None:
        raise HTTPException(404, "Апелляция не найдена")
    if a.status in (AppealStatus.ACCEPTED, AppealStatus.REJECTED):
        raise HTTPException(409, "Решение уже принято")
    if data.status in (AppealStatus.ACCEPTED, AppealStatus.REJECTED) and not data.decision:
        raise HTTPException(422, "Укажите обоснование решения")
    a.status = data.status
    if data.status == AppealStatus.UNDER_REVIEW:
        _event(db, "appeal", a.id, a.status.value, p.id, data.decision or "Принята к рассмотрению")
    else:
        a.decision = data.decision
        a.decided_at = utcnow()
        _event(db, "appeal", a.id, a.status.value, p.id, data.decision)
        if data.status == AppealStatus.ACCEPTED and data.new_score is not None:
            res = db.scalar(select(Result).where(Result.competition_id == a.competition_id, Result.athlete_id == a.athlete_id))
            old = res.score
            res.score, res.place_manual = data.new_score, False
            a.new_score = data.new_score
            _recalc_and_notify(db, a.competition, f"по апелляции изменены баллы {a.athlete.full_name}")
            log_action(db, p, "correction", "protocol", a.competition_id, appeal_id=a.id, old=old, new=data.new_score)
    notify(db, a.athlete_id, "Апелляция", f"«{a.competition.title}»: статус «{a.status.value}»", "/athlete/appeals")
    if a.submitted_by_id != a.athlete_id:
        notify(db, a.submitted_by_id, "Апелляция", f"«{a.competition.title}»: статус «{a.status.value}»", None)
    log_action(db, p, "status", "appeal", a.id, new=a.status.value)
    db.commit()
    bg.add_task(deliver_in_background)
    return appeal_out(db, a)


def violation_out(db: DB, v: Violation) -> dict:
    return {
        "id": v.id, "competition_id": v.competition_id, "competition": v.competition.title,
        "athlete_id": v.athlete_id, "athlete": v.athlete.full_name, "report": v.report,
        "reported_by": db.get(User, v.reported_by_id).full_name, "decision": v.decision.value,
        "decision_text": v.decision_text, "disqualification_until": v.disqualification_until,
        "created_at": v.created_at.isoformat(), **_case_details(db, "violation", v.id),
    }


@router.post("/violations", status_code=201)
async def report_violation(
    db: DB, p: CurrentPrincipal,
    competition_id: int = Form(...), athlete_id: int = Form(...), report: str = Form(..., min_length=10),
    files: list[UploadFile] = File(default=[]),
):
    comp = db.get(Competition, competition_id)
    if comp is None:
        raise HTTPException(404, "Соревнование не найдено")
    in_panel = db.scalar(select(JudgeAssignment.id).where(JudgeAssignment.competition_id == comp.id, JudgeAssignment.judge_id == p.id))
    if p.role != Role.ADMIN and not (p.role == Role.JUDGE and in_panel):
        raise HTTPException(403, "Рапорт подаёт федерация или судья коллегии")
    v = Violation(competition_id=comp.id, athlete_id=athlete_id, reported_by_id=p.id, report=report)
    db.add(v)
    db.flush()
    await _attach(db, "violation", v.id, files, p.id)
    _event(db, "violation", v.id, "reported", p.id, "Рапорт зарегистрирован")
    notify_role(db, Role.ADMIN, "Рапорт о нарушении", f"«{comp.title}»", "/admin/appeals")
    log_action(db, p, "create", "violation", v.id)
    db.commit()
    return violation_out(db, v)


@router.get("/violations")
def list_violations(db: DB, p: CurrentPrincipal):
    stmt = select(Violation).order_by(Violation.id.desc())
    if p.role == Role.ATHLETE:
        stmt = stmt.where(Violation.athlete_id == p.id)
    elif p.role == Role.JUDGE:
        stmt = stmt.where(Violation.reported_by_id == p.id)
    elif p.role not in (Role.ADMIN, Role.MINISTRY):
        raise HTTPException(403, "Нет доступа")
    return [violation_out(db, v) for v in db.scalars(stmt)]


class ViolationDecisionIn(BaseModel):
    decision: ViolationDecision
    decision_text: str
    disqualification_until: date | None = None


@router.post("/violations/{vid}/decision")
def decide_violation(vid: int, data: ViolationDecisionIn, db: DB, p: AdminPrincipal, bg: BackgroundTasks):
    v = db.get(Violation, vid)
    if v is None:
        raise HTTPException(404, "Рапорт не найден")
    if v.decision != ViolationDecision.PENDING:
        raise HTTPException(409, "Решение уже принято")
    if data.decision == ViolationDecision.PENDING:
        raise HTTPException(422, "Выберите решение")
    v.decision, v.decision_text, v.decided_at = data.decision, data.decision_text, utcnow()
    res = db.scalar(select(Result).where(Result.competition_id == v.competition_id, Result.athlete_id == v.athlete_id))
    if data.decision in (ViolationDecision.TECHNICAL_DEFEAT, ViolationDecision.DISQUALIFICATION) and res:
        res.score, res.place, res.place_manual = None, None, False
        res.note = "Техническое поражение" if data.decision == ViolationDecision.TECHNICAL_DEFEAT else "Дисквалифицирован"
        _recalc_and_notify(db, v.competition, f"решение по нарушению ({res.note.lower()})")
        log_action(db, p, "correction", "protocol", v.competition_id, violation_id=v.id)
    if data.decision == ViolationDecision.DISQUALIFICATION:
        v.disqualification_until = data.disqualification_until
        db.add(Disqualification(athlete_id=v.athlete_id, date_from=date.today(), date_to=data.disqualification_until,
                                reason=f"Решение по рапорту № {v.id}: {data.decision_text}"))
    _event(db, "violation", v.id, data.decision.value, p.id, data.decision_text)
    notify(db, v.athlete_id, "Решение по нарушению", f"«{v.competition.title}»: {data.decision_text}", None)
    log_action(db, p, "status", "violation", v.id, decision=data.decision.value)
    db.commit()
    bg.add_task(deliver_in_background)
    return violation_out(db, v)


@router.get("/evidence/{eid}")
def download_evidence(eid: int, db: DB, p: CurrentPrincipal):
    e = db.get(Evidence, eid)
    if e is None:
        raise HTTPException(404, "Файл не найден")
    case = db.get(Appeal if e.subject == "appeal" else Violation, e.subject_id)
    allowed = p.role in (Role.ADMIN, Role.MINISTRY) or p.id in (
        e.uploaded_by_id, case.athlete_id, getattr(case, "submitted_by_id", None), getattr(case, "reported_by_id", None))
    if not allowed:
        raise HTTPException(403, "Нет доступа")
    return FileResponse(folder("evidence") / e.stored_as, filename=e.filename)




def prize_candidates(db: DB, year: int) -> list[dict]:
    rules = {(r.rating_level, r.place): r.amount for r in db.scalars(select(PrizeRule))}
    comps = db.scalars(select(Competition).where(
        Competition.status == CompetitionStatus.COMPLETED,
        Competition.start_date.between(date(year, 1, 1), date(year, 12, 31)),
    ).order_by(Competition.start_date)).all()
    by_athlete: dict[int, dict] = {}
    check_on = min(date.today(), date(year, 12, 31))
    for c in comps:
        for r in c.results:
            amount = rules.get((c.rating_level, r.place)) if r.place else None
            if not amount:
                continue
            dq = [d for d in db.scalars(select(Disqualification).where(Disqualification.athlete_id == r.athlete_id)) if d.is_active_on(check_on)]
            if dq:
                continue
            u = r.athlete
            item = by_athlete.setdefault(u.id, {
                "athlete_id": u.id, "full_name": u.full_name,
                "birth_date": u.birth_date.isoformat() if u.birth_date else None,
                "municipality": u.municipality.name if u.municipality else None,
                "passport": docs.mask("4510123456"), "inn": docs.mask("561234567890"), "snils": docs.mask("12345678901"),
                "results": [], "amount": 0,
            })
            item["results"].append({"competition": c.title, "date": c.start_date.isoformat(), "level": LEVEL_RU[c.rating_level],
                                    "place": r.place, "amount": amount})
            item["amount"] += amount
    return sorted(by_athlete.values(), key=lambda i: -i["amount"])


@router.get("/prizes/candidates")
def get_prize_candidates(db: DB, _: ReaderPrincipal, year: int | None = None):
    year = year or date.today().year
    items = prize_candidates(db, year)
    return {"year": year, "deadline": date(year + 1, 2, 1), "items": items, "total": sum(i["amount"] for i in items)}


class PrizeRuleIO(BaseModel):
    rating_level: RatingLevel
    place: int
    amount: int


@router.get("/admin/settings/prizes", response_model=list[PrizeRuleIO])
def get_prize_rules(db: DB, _: ReaderPrincipal):
    return [PrizeRuleIO(rating_level=r.rating_level, place=r.place, amount=r.amount)
            for r in db.scalars(select(PrizeRule).order_by(PrizeRule.rating_level, PrizeRule.place))]


@router.put("/admin/settings/prizes", response_model=list[PrizeRuleIO])
def set_prize_rules(rules: list[PrizeRuleIO], db: DB, p: AdminPrincipal):
    db.execute(delete(PrizeRule))
    for r in rules:
        db.add(PrizeRule(**r.model_dump()))
    log_action(db, p, "update", "prize_rules", None, count=len(rules))
    db.commit()
    return get_prize_rules(db, p)


def petition_out(pt: PrizePetition) -> dict:
    return {"id": pt.id, "year": pt.year, "status": pt.status.value, "total_amount": pt.total_amount,
            "items": pt.items, "created_at": pt.created_at.isoformat(),
            "sent_at": pt.sent_at.isoformat() if pt.sent_at else None,
            "awarded_at": pt.awarded_at.isoformat() if pt.awarded_at else None}


class PetitionIn(BaseModel):
    year: int
    athlete_ids: list[int] | None = None


@router.post("/prizes/petitions", status_code=201)
def create_petition(data: PetitionIn, db: DB, p: AdminPrincipal):
    items = prize_candidates(db, data.year)
    if data.athlete_ids is not None:
        items = [i for i in items if i["athlete_id"] in set(data.athlete_ids)]
    if not items:
        raise HTTPException(409, "Нет претендентов на премию")
    pt = PrizePetition(year=data.year, items=items, total_amount=sum(i["amount"] for i in items))
    db.add(pt)
    db.flush()
    log_action(db, p, "create", "prize_petition", pt.id, total=pt.total_amount)
    db.commit()
    return petition_out(pt)


@router.get("/prizes/petitions")
def list_petitions(db: DB, _: ReaderPrincipal):
    return [petition_out(pt) for pt in db.scalars(select(PrizePetition).order_by(PrizePetition.id.desc()))]


@router.post("/prizes/petitions/{pid}/status")
def petition_status(pid: int, db: DB, p: AdminPrincipal, status: Annotated[PetitionStatus, Body(embed=True)]):
    pt = db.get(PrizePetition, pid)
    if pt is None:
        raise HTTPException(404, "Ходатайство не найдено")
    order = list(PetitionStatus)
    if order.index(status) != order.index(pt.status) + 1:
        raise HTTPException(409, "Статусы меняются по порядку: подготовлено → направлено → премия назначена")
    pt.status = status
    if status == PetitionStatus.SENT:
        pt.sent_at = utcnow()
    if status == PetitionStatus.AWARDED:
        pt.awarded_at = utcnow()
        for i in pt.items:
            notify(db, i["athlete_id"], "Назначена премия Губернатора", f"Сумма: {i['amount']:,} ₽".replace(",", " "), None)
    log_action(db, p, "status", "prize_petition", pt.id, new=status.value)
    db.commit()
    return petition_out(pt)


@router.get("/prizes/petitions/{pid}.docx")
def petition_docx(pid: int, db: DB, _: ReaderPrincipal):
    pt = db.get(PrizePetition, pid)
    if pt is None:
        raise HTTPException(404, "Ходатайство не найдено")
    doc = docs.new_doc("ХОДАТАЙСТВО", [
        "Министру физической культуры и спорта Оренбургской области",
        "",
        f"В соответствии с Указом Губернатора Оренбургской области от 04.03.2011 № 93-ук региональное отделение "
        f"Федерации спортивного программирования ходатайствует о назначении премий спортсменам по итогам {pt.year} года "
        f"за высокие спортивные результаты согласно приложению.",
        f"Общая сумма: {pt.total_amount:,} ₽.".replace(",", " "),
        "",
        "Приложение 1. Список претендентов.",
    ])
    docs.add_table(doc, ["№", "ФИО", "Дата рождения", "Паспорт", "ИНН", "СНИЛС", "Результаты", "Сумма, ₽"], [
        (n, i["full_name"], i["birth_date"], i["passport"], i["inn"], i["snils"],
         "; ".join(f"{r['competition']} - {r['place']} место" for r in i["results"]), i["amount"])
        for n, i in enumerate(pt.items, 1)
    ])
    doc.add_paragraph("Приложение 2. Копии протоколов соревнований (выгружаются из платформы).")
    doc.add_paragraph("Персональные данные в демонстрационной версии замаскированы.")
    docs.signature(doc, "Председатель регионального отделения")
    return docs.file_response(docs.docx_bytes(doc), f"governor_prize_{pt.year}_{pt.id}.docx")
