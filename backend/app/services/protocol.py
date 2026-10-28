
from datetime import date
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import (
    AthleteRank,
    Application,
    ApplicationStatus,
    Competition,
    JudgeQualification,
    Result,
)
from ..schemas import JudgeRow, ProtocolRow
from . import evsk

LEVEL_RU = {"regional": "региональное", "interregional": "межрегиональное", "all_russian": "всероссийское"}
FORMAT_RU = {"offline": "очный", "online": "онлайн", "mixed": "смешанный"}


def assign_places(results: list[Result]) -> None:
    auto = sorted((r for r in results if not r.place_manual and r.score is not None), key=lambda r: -r.score)
    prev_score, prev_place = None, 0
    for i, r in enumerate(auto, start=1):
        if r.score != prev_score:
            prev_place, prev_score = i, r.score
        r.place = prev_place
    for r in results:
        if not r.place_manual and r.score is None:
            r.place = None


def current_rank_name(db: Session, athlete_id: int, discipline_id: int, on: date) -> str | None:
    ranks = db.scalars(
        select(AthleteRank).where(
            AthleteRank.athlete_id == athlete_id, AthleteRank.discipline_id == discipline_id
        )
    ).all()
    valid = [r for r in ranks if r.valid_until is None or r.valid_until >= on]
    return max(valid, key=lambda r: r.rank.level).rank.name if valid else None


def build_rows(db: Session, comp: Competition) -> list[ProtocolRow]:
    approved = db.scalars(
        select(Application).where(
            Application.competition_id == comp.id, Application.status == ApplicationStatus.APPROVED
        )
    ).all()
    results = {r.athlete_id: r for r in comp.results}
    participants = sum(1 for r in comp.results if r.place)
    rows: list[ProtocolRow] = []
    for app in approved:
        a = app.athlete
        res = results.get(a.id)
        prof = a.athlete
        rows.append(
            ProtocolRow(
                place=res.place if res else None,
                athlete_id=a.id,
                full_name=a.full_name,
                birth_year=a.birth_date.year if a.birth_date else None,
                municipality=a.municipality.name if a.municipality else None,
                organization=(prof.organization.short_name or prof.organization.name)
                if prof and prof.organization
                else "без организации",
                rank=current_rank_name(db, a.id, comp.discipline_id, comp.start_date),
                score=res.score if res else None,
                note=res.note if res else None,
                result_id=res.id if res else None,
                evsk_hint=_hint_name(db, a.id, comp, res, participants),
            )
        )
    rows.sort(key=lambda r: (r.place is None, r.place or 0, r.full_name))
    return rows


def _hint_name(db: Session, athlete_id: int, comp: Competition, res: Result | None, participants: int) -> str | None:
    if res is None or not res.place:
        return None
    r = evsk.hint(db, athlete_id, comp.discipline_id, comp.rating_level, res.place, participants, comp.start_date)
    return r.name if r else None


def build_judges(db: Session, comp: Competition) -> list[JudgeRow]:
    out = []
    for ja in comp.judges:
        q = db.scalars(
            select(JudgeQualification)
            .where(JudgeQualification.judge_id == ja.judge_id)
            .order_by(JudgeQualification.assigned_at.desc())
        ).first()
        out.append(JudgeRow(full_name=ja.judge.full_name, position=ja.position.name, category=q.category.name if q else None))
    return out


def _header_lines(comp: Competition, number: str | None) -> list[str]:
    dates = comp.start_date.strftime("%d.%m.%Y")
    if comp.end_date != comp.start_date:
        dates += " – " + comp.end_date.strftime("%d.%m.%Y")
    lines = [
        "Региональное отделение Федерации спортивного программирования Оренбургской области",
        f"ПРОТОКОЛ{f' № {number}' if number else ''} результатов",
        comp.title,
        f"Дисциплина: {comp.discipline.name}",
        f"Статус: {LEVEL_RU[comp.level.value]}; тип: {comp.type.name}; возрастная группа: {comp.age_group.name}",
        f"Сроки: {dates}; формат: {FORMAT_RU[comp.format.value]}" + (f"; место: {comp.venue}" if comp.venue else ""),
    ]
    if comp.calendar_plan_number:
        plan = "ЕКП Минспорта России" if comp.calendar_plan.value == "ekp" else "календарный план Оренбургской области"
        lines.append(f"Включено в {plan}, № {comp.calendar_plan_number}")
    return lines


COLUMNS = ["Место", "ФИО", "Год рожд.", "Муниципалитет", "Организация", "Разряд", "Баллы", "Примечание"]


def _row_values(r: ProtocolRow) -> list:
    return [r.place or "-", r.full_name, r.birth_year, r.municipality, r.organization, r.rank or "б/р", r.score, r.note]


def _signature_lines(judges: list[JudgeRow]) -> list[str]:
    lead = [j for j in judges if j.position.lower().startswith("главн")]
    return [f"{j.position}: ________________ / {j.full_name} /" for j in lead]


def export_xlsx(comp: Competition, number: str | None, rows: list[ProtocolRow], judges: list[JudgeRow]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Протокол"
    header = _header_lines(comp, number)
    for i, line in enumerate(header, start=1):
        ws.cell(row=i, column=1, value=line).font = Font(bold=i in (2, 3))
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=len(COLUMNS))
        ws.cell(row=i, column=1).alignment = Alignment(horizontal="center" if i <= 3 else "left", wrap_text=True)
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    start = len(header) + 2
    for c, title in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=start, column=c, value=title)
        cell.font = Font(bold=True)
        cell.border = border
    for i, r in enumerate(rows, start=start + 1):
        for c, v in enumerate(_row_values(r), start=1):
            ws.cell(row=i, column=c, value=v).border = border
    for col, width in zip("ABCDEFGH", (8, 34, 10, 22, 30, 10, 10, 24)):
        ws.column_dimensions[col].width = width
    line = start + len(rows) + 2
    for s in _signature_lines(judges):
        ws.cell(row=line, column=1, value=s)
        line += 1
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_docx(comp: Competition, number: str | None, rows: list[ProtocolRow], judges: list[JudgeRow]) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    for i, line in enumerate(_header_lines(comp, number)):
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.bold = i in (1, 2)
        if i <= 2:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = doc.add_table(rows=1, cols=len(COLUMNS))
    table.style = "Table Grid"
    for c, title in enumerate(COLUMNS):
        table.rows[0].cells[c].text = title
        table.rows[0].cells[c].paragraphs[0].runs[0].bold = True
    for r in rows:
        cells = table.add_row().cells
        for c, v in enumerate(_row_values(r)):
            cells[c].text = "" if v is None else str(v)
    doc.add_paragraph()
    for s in _signature_lines(judges):
        doc.add_paragraph(s)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
