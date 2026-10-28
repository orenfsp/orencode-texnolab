
from collections.abc import Iterable, Sequence
from datetime import date
from io import BytesIO
from typing import Any
from urllib.parse import quote

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Border, Font, Side

FEDERATION = "Региональное отделение Федерации спортивного программирования Оренбургской области"
MIME = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def new_doc(title: str, subtitle_lines: Iterable[str] = ()) -> DocxDocument:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    p = doc.add_paragraph(FEDERATION)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(14)
    for line in subtitle_lines:
        doc.add_paragraph(line)
    return doc


def add_table(doc: DocxDocument, columns: Sequence[str], rows: Iterable[Sequence[Any]]) -> None:
    t = doc.add_table(rows=1, cols=len(columns))
    t.style = "Table Grid"
    for i, c in enumerate(columns):
        t.rows[0].cells[i].text = c
        t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = "" if v is None else str(v)


def signature(doc: DocxDocument, position: str, name: str = "") -> None:
    doc.add_paragraph()
    doc.add_paragraph(f"{position} ________________ / {name or '                    '} /")
    doc.add_paragraph(f"«___» ____________ {date.today().year} г.")


def docx_bytes(doc: DocxDocument) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def xlsx_bytes(sheets: dict[str, tuple[Sequence[str], Iterable[Sequence[Any]]]], title: str | None = None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for name, (columns, rows) in sheets.items():
        ws = wb.create_sheet(name[:31])
        start = 1
        if title:
            ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=12)
            start = 3
        for c, h in enumerate(columns, start=1):
            cell = ws.cell(row=start, column=c, value=h)
            cell.font = Font(bold=True)
            cell.border = border
            ws.column_dimensions[cell.column_letter].width = max(12, min(48, len(str(h)) + 6))
        for r, row in enumerate(rows, start=start + 1):
            for c, v in enumerate(row, start=1):
                ws.cell(row=r, column=c, value=v).border = border
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def file_response(content: bytes, filename: str) -> Response:
    ext = filename.rsplit(".", 1)[-1]
    return Response(
        content,
        media_type=MIME.get(ext, "application/octet-stream"),
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


def mask(value: str | None, keep: int = 2) -> str:
    if not value:
        return "•••••••"
    return "•" * max(0, len(value) - keep) + value[-keep:]
